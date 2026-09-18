import argparse
import csv
import os
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from cambridge.registration import register_orders as register_candidates_batch
from update_status import main as update_xnet_status
import sheets

DATA_DIR = Path(os.getenv("DATA_DIR", "/var/data"))
CSV_PATH = DATA_DIR / "orders.csv"
RESULTS_CSV_PATH = DATA_DIR / "cambridge_results.csv"
REPORTS_DIR = DATA_DIR / "reports"
DATABASE_XLSX_PATH = DATA_DIR / "cambridge_registrations.xlsx"

RESULT_FIELDS = [
    "timestamp",
    "order_number",
    "email",
    "password",
    "session_name",
    "success",
    "manual_review_reason",
    "confirmation",
]

# Columns of the persistent "database" of validated (successful)
# registrations only — one row per order_number, updated in place on rerun.
DATABASE_FIELDS = [
    "order_number",
    "surname",
    "name",
    "email",
    "password",
    "date_of_birth",
    "gender",
    "nationality",
    "id_number",
    "linguaskill_type",
    "exam_type",
    "skills_code",
    "exam_date",
    "exam_hour",
    "session_name",
    "confirmation",
    "registered_at",
]


def register_orders(df: pd.DataFrame, headless: bool = True) -> dict:
    """Register every order in `df` against Cambridge/Metrica.

    Logs into Cambridge/Metrica once for the whole batch (cambridge.
    registration.register_orders) rather than per order — faster, and a
    failure on any one order (session creation, credits, a form error...)
    is recorded and the loop moves on to the next order; it never aborts
    the whole run and never needs restarting from scratch.

    For each order: writes the result back to X-Net (update_status.py —
    this is where the password is durably recorded for future reuse),
    appends a row to the local results CSV, updates the local Excel
    "database" of validated registrations, and writes a human-readable
    report for the run.

    Returns the same {order_number: {...}} dict passed to update_status.py.
    """
    timestamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
    orders = [row.to_dict() for _, row in df.iterrows()]
    results = register_candidates_batch(orders, headless=headless)

    rows = []
    db_rows = []
    for order in orders:
        order_number = order.get("order_number")
        result = results[order_number]
        rows.append(
            {
                "timestamp": timestamp,
                "order_number": order_number,
                "email": result.get("email"),
                "password": result.get("password"),
                "session_name": result.get("session_name", order.get("session_name")),
                "success": result.get("success"),
                "manual_review_reason": result.get("manual_review_reason", ""),
                "confirmation": result.get("confirmation", ""),
            }
        )
        if result.get("success"):
            db_rows.append(
                {
                    "order_number": order_number,
                    "surname": order.get("surname"),
                    "name": order.get("name"),
                    "email": result.get("email"),
                    "password": result.get("password"),
                    "date_of_birth": order.get("date_of_birth"),
                    "gender": order.get("gender"),
                    "nationality": order.get("nationality"),
                    "id_number": order.get("id_number"),
                    "linguaskill_type": order.get("linguaskill_type"),
                    "exam_type": order.get("exam_type"),
                    "skills_code": order.get("skills_code"),
                    "exam_date": order.get("exam_date"),
                    "exam_hour": order.get("exam_hour"),
                    "session_name": result.get("session_name", order.get("session_name")),
                    "confirmation": result.get("confirmation"),
                    "registered_at": timestamp,
                }
            )

    if rows:
        _append_results_csv(rows)
        _write_report(rows, timestamp)

    if db_rows:
        # Google Sheets is the durable, team-shared copy (survives Render
        # restarts with no persistent disk needed). The local xlsx is a
        # best-effort convenience copy for local runs — its failure (or
        # Sheets' unavailability) never aborts the batch.
        try:
            sheets.upsert_registrations(db_rows, DATABASE_FIELDS)
        except Exception as exc:
            print(f"[WARN] could not update Google Sheet: {exc}")
        try:
            _update_database_xlsx(db_rows)
        except Exception as exc:
            print(f"[WARN] could not update local xlsx database: {exc}")

    if results:
        update_xnet_status(results, headless=headless)

    return results


def _update_database_xlsx(new_rows: list[dict]):
    """Upsert validated registrations into the local Excel database.

    One row per order_number: a rerun of the same order (e.g. a retry that
    now succeeds) replaces its previous row rather than duplicating it.
    """
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    new_df = pd.DataFrame(new_rows, columns=DATABASE_FIELDS)

    if DATABASE_XLSX_PATH.exists():
        existing_df = pd.read_excel(DATABASE_XLSX_PATH, dtype=str)
        combined = pd.concat([existing_df, new_df], ignore_index=True)
    else:
        combined = new_df

    combined = combined.drop_duplicates(subset="order_number", keep="last")
    combined = combined.sort_values("registered_at").reset_index(drop=True)

    combined.to_excel(DATABASE_XLSX_PATH, index=False, sheet_name="Registrations")
    _format_xlsx(DATABASE_XLSX_PATH, combined)


def _format_xlsx(path: Path, df: pd.DataFrame):
    """Bold header row and auto-size columns for readability."""
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, column in enumerate(df.columns, start=1):
        width = max(len(str(column)), df[column].astype(str).map(len).max() if len(df) else 0)
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 40)
    ws.freeze_panes = "A2"
    wb.save(path)


def _append_results_csv(rows: list[dict]):
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    file_exists = RESULTS_CSV_PATH.exists()
    with open(RESULTS_CSV_PATH, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=RESULT_FIELDS)
        if not file_exists:
            writer.writeheader()
        writer.writerows(rows)


def _write_report(rows: list[dict], timestamp: str) -> Path:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    safe_ts = timestamp.replace(":", "-")
    path = REPORTS_DIR / f"cambridge_run_{safe_ts}.md"

    ok_count = sum(1 for r in rows if r["success"])
    lines = [
        f"# Cambridge registration run — {timestamp}",
        "",
        f"{ok_count}/{len(rows)} orders registered successfully.",
        "",
    ]
    for r in rows:
        confirmation = r.get("confirmation") or ""
        if not r["success"]:
            status = "MANUAL REVIEW"
        elif "skipped duplicate registration" in confirmation:
            status = "SUCCESS (candidat déjà trouvé dans la session — pas réinscrit)"
        else:
            status = "SUCCESS"
        lines.append(f"## Order {r['order_number']} — {status}")
        lines.append(f"- Email: {r['email']}")
        lines.append(f"- Session: {r['session_name']}")
        if r["success"]:
            lines.append(f"- Confirmation: {confirmation}")
        else:
            lines.append(f"- Reason: {r['manual_review_reason']}")
        lines.append("")

    path.write_text("\n".join(lines), encoding="utf-8")
    return path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Register orders against Cambridge/Metrica."
    )
    parser.add_argument("--headed", action="store_true", help="Run with a visible browser window.")
    parser.add_argument(
        "--csv",
        default=None,
        help=f"Path to an orders CSV to process instead of fetching fresh from Victoria/X-Net (default: fetch fresh via runner.main()).",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Only process the first N orders — use this for a first real test.",
    )
    parser.add_argument(
        "--order",
        default=None,
        help="Only process this one order_number — safest option for a first real test.",
    )
    args = parser.parse_args()

    if args.csv:
        orders_df = pd.read_csv(args.csv)
        print(f"[INFO] {len(orders_df)} orders loaded from {args.csv}")
    else:
        from runner import main as fetch_victoria_orders

        orders_df = fetch_victoria_orders(headless=not args.headed)
        if orders_df is None:
            # runner.main() prints its own "[ERROR] ..." and returns None
            # on failure (e.g. a stale Victoria login or a transient
            # navigation error) instead of raising — nothing to register.
            print("[ERROR] Victoria/X-Net order fetch failed (see error above). Nothing to register.")
            raise SystemExit(1)
        print(f"[INFO] {len(orders_df)} orders fetched fresh from Victoria/X-Net")
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        orders_df.to_csv(CSV_PATH, index=False)

    if args.order:
        orders_df = orders_df[orders_df["order_number"] == args.order]
        print(f"[INFO] filtered to order {args.order!r}: {len(orders_df)} row(s)")
    elif args.limit:
        orders_df = orders_df.head(args.limit)
        print(f"[INFO] limited to first {args.limit} order(s)")

    if len(orders_df) == 0:
        print("[INFO] nothing to register.")
    else:
        outcome = register_orders(orders_df, headless=not args.headed)
        print(f"[INFO] {sum(1 for r in outcome.values() if r['success'])}/{len(outcome)} succeeded")
        for order_number, result in outcome.items():
            if not result.get("success"):
                print(f"[MANUAL REVIEW] {order_number}: {result.get('manual_review_reason')}")
